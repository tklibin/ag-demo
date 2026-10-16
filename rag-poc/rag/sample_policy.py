"""A fictional emergency-department claims policy, written as both PDF and DOCX, for testing.

"Example Health Assurance" is invented. The content mirrors the rules simulated in the CDI and
funding POCs (CTAS triage levels, prior authorisation for low-acuity CT, high-cost review,
clawbacks) so the three POCs tell one story. The PDF gets a running header and page footer on
every page, which the parser is expected to strip.
"""

from __future__ import annotations

import html
from pathlib import Path

import pymupdf

TITLE = "Emergency Department Claims Policy"
ISSUER = "Example Health Assurance (fictional insurer)"
VERSION = "Version 3.0, effective 1 January 2027. Synthetic sample document for testing only."

# (level, heading, paragraphs). Level 1 = numbered section, level 2 = clause.
SECTIONS: list[tuple[int, str, list[str]]] = [
    (1, "1 Purpose and scope", []),
    (2, "1.1 Purpose", [
        "This policy sets out how Example Health Assurance (the Insurer) reviews, approves and pays claims "
        "for services delivered to its members in hospital emergency departments (ED)."]),
    (2, "1.2 Scope", [
        "The policy applies to all network hospitals. It takes effect on 1 January 2027 and replaces "
        "version 2.3 of the policy for all services delivered on or after that date."]),
    (2, "1.3 Services outside this policy", [
        "Inpatient care after the first 24 hours of an admission from the ED is governed by the Inpatient "
        "Claims Policy. Outpatient pharmacy claims are governed by the Pharmacy Benefit Policy."]),

    (1, "2 Definitions", []),
    (2, "2.1 Emergency", [
        "An emergency is a condition that, without immediate care, could reasonably be expected to place the "
        "member's life or long-term health in serious danger. For this policy, patients triaged CTAS 1, 2 or 3 "
        "are treated as emergencies."]),
    (2, "2.2 CTAS", [
        "CTAS means the Canadian Triage and Acuity Scale, recorded at triage, from level 1 (resuscitation) to "
        "level 5 (non-urgent). The level recorded at first triage determines the authorisation rules; a later "
        "re-triage to a more urgent level applies from the time it is recorded."]),
    (2, "2.3 Clean claim", [
        "A clean claim contains every mandatory field and attachment, relates to an eligible member and passes "
        "the Insurer's automated validation without manual correction."]),

    (1, "3 Eligibility verification", []),
    (2, "3.1 Verification before discharge", [
        "The hospital must verify the member's eligibility through the Insurer's eligibility service on the "
        "date of service, before the patient is discharged from the ED where possible."]),
    (2, "3.2 Unverified eligibility", [
        "Claims for members whose eligibility was not verified will be rejected, unless the hospital provides "
        "proof of a successful verification within 10 working days of the rejection notice."]),
    (2, "3.3 Unconscious or unidentified patients", [
        "For patients who arrive unconscious or cannot be identified, eligibility may be verified up to 72 "
        "hours after arrival."]),

    (1, "4 Prior authorisation", []),
    (2, "4.1 Emergency services", [
        "No prior authorisation is required for any service delivered to a patient triaged CTAS 1, 2 or 3."]),
    (2, "4.2 Non-emergency services", [
        "For patients triaged CTAS 4 or 5, prior authorisation is required before CT imaging, MRI, minor "
        "procedures and observation stays longer than 6 hours. Laboratory tests, plain X-rays and ultrasound "
        "do not require prior authorisation. Annex A lists every service and the rule that applies."]),
    (2, "4.3 Retrospective authorisation", [
        "If the condition of a CTAS 4 or 5 patient deteriorates and a service requiring authorisation is "
        "delivered without it, the hospital may request retrospective authorisation within 24 hours of the "
        "service. Requests received after 24 hours will be declined and the related line items will not be paid."]),
    (2, "4.4 Response times", [
        "The Insurer responds to prior authorisation requests for ED services within 60 minutes. A request that "
        "receives no response within 60 minutes is deemed approved, and the hospital must record the request "
        "reference on the claim."]),

    (1, "5 Documentation requirements", []),
    (2, "5.1 Signature", [
        "Every clinical note must be signed electronically by the treating physician before the claim is "
        "submitted. Claims with unsigned notes will be returned as incomplete."]),
    (2, "5.2 Imaging", [
        "Requests for CT or MRI must document the clinical indication and the relevant examination findings. "
        "For CTAS 4 or 5 patients the note must also explain why the investigation could not be deferred to an "
        "outpatient appointment."]),
    (2, "5.3 Laterality", [
        "For conditions of paired organs or limbs, the side affected (left, right or bilateral) must be "
        "documented in the note and reflected in the diagnosis code."]),
    (2, "5.4 Observation", [
        "Observation stays must record the start and end time and a documented reassessment at least every "
        "4 hours."]),

    (1, "6 Coding standards", []),
    (2, "6.1 Code sets", [
        "Diagnoses must be coded in ICD-10-AM and procedures in the national billing code set in force on the "
        "date of service."]),
    (2, "6.2 Unspecified codes", [
        "Unspecified diagnosis codes are accepted only when the documentation does not support a more specific "
        "code. Where the documentation supports a more specific code, the affected line items may be reduced by "
        "10 percent."]),

    (1, "7 Claim submission", []),
    (2, "7.1 Submission deadline", [
        "Claims must be submitted within 30 days of the patient's discharge from the ED.",
        "Claims submitted between 31 and 90 days after discharge are accepted with a 5 percent late-submission "
        "reduction. Claims submitted more than 90 days after discharge will be rejected as late."]),
    (2, "7.2 Resubmission", [
        "A rejected claim may be corrected and resubmitted once, within 30 days of the rejection notice."]),
    (2, "7.3 Attachments for high-value claims", [
        "Claims with a billed amount above SAR 5,000 must include the ED discharge summary."]),

    (1, "8 Adjudication and payment", []),
    (2, "8.1 Adjudication", [
        "The Insurer adjudicates clean claims within 30 days of receipt."]),
    (2, "8.2 Payment", [
        "Approved claims are paid within 45 days of the adjudication decision, in the payment run on the 15th "
        "of each month."]),
    (2, "8.3 Late payment", [
        "If an approved claim is paid more than 15 days after the deadline in clause 8.2, the Insurer adds a "
        "late-payment charge of 1 percent per month on the unpaid amount."]),

    (1, "9 Reductions and high-cost review", []),
    (2, "9.1 Line-item reductions", [
        "Line items that the documentation does not support are reduced or denied. The remittance advice lists "
        "the reason code for each reduced line."]),
    (2, "9.2 High-cost review", [
        "Claims with a billed amount above SAR 4,000 undergo a line-by-line medical review before payment. The "
        "review may extend the adjudication period in clause 8.1 by up to 20 days."]),
    (2, "9.3 Observation beyond 6 hours", [
        "For CTAS 4 or 5 patients, observation beyond 6 hours without prior authorisation is paid at the "
        "6-hour rate."]),

    (1, "10 Audit and recoupment", []),
    (2, "10.1 Audit window", [
        "The Insurer may audit any paid claim within 12 months of the payment date."]),
    (2, "10.2 Notice and records", [
        "The hospital receives 30 days' written notice of an audit and must provide the requested records "
        "within 15 working days of the audit start date."]),
    (2, "10.3 Recovery of overpayments", [
        "Amounts found to be overpaid are recovered by offset against future payments to the hospital, after "
        "the appeal period in section 11 has ended without an appeal or after a final appeal decision."]),

    (1, "11 Appeals", []),
    (2, "11.1 Right to appeal", [
        "A hospital may appeal a rejection, a reduction or a recoupment within 30 days of the decision."]),
    (2, "11.2 Appeal levels", [
        "A first-level appeal is decided by the claims review team within 15 working days. The hospital may "
        "escalate to a second-level appeal, decided by a panel chaired by the medical director within 30 days. "
        "The second-level decision is final."]),
    (2, "11.3 New information", [
        "Appeals must include new information or documentation. An appeal that only restates the original "
        "claim will be dismissed."]),

    (1, "12 Exclusions", []),
    (2, "12.1 Excluded services", [
        "The following ED services are not covered: cosmetic procedures; examinations requested only for "
        "administrative purposes, such as fitness-to-work certificates; and non-emergency dental treatment."]),
]

ANNEX_A = ("Annex A – Prior authorisation list for CTAS 4 and 5", ["Service", "Authorisation", "Notes"], [
    ["Laboratory tests", "Not required", ""],
    ["Plain X-ray", "Not required", ""],
    ["Ultrasound", "Not required", ""],
    ["CT, any region", "Required", "Clause 5.2 documentation applies"],
    ["MRI, any region", "Required", "Clause 5.2 documentation applies"],
    ["Minor procedures", "Required", "Except simple wound closure under 5 cm, which needs no authorisation"],
    ["Observation up to 6 hours", "Not required", "Clause 5.4 documentation applies"],
    ["Observation beyond 6 hours", "Required", "Otherwise paid at the 6-hour rate (clause 9.3)"],
])


def write_pdf(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    body = [f"<h1>{html.escape(TITLE)}</h1>", f"<p><b>{html.escape(ISSUER)}</b></p>", f"<p>{html.escape(VERSION)}</p>"]
    for level, heading, paragraphs in SECTIONS:
        body.append(f"<h{level + 1}>{html.escape(heading)}</h{level + 1}>")
        body += [f"<p>{html.escape(p)}</p>" for p in paragraphs]
    title, header, rows = ANNEX_A
    body.append(f"<h2>{html.escape(title)}</h2><table>")
    body.append("<tr>" + "".join(f"<th>{html.escape(h)}</th>" for h in header) + "</tr>")
    for row in rows:
        body.append("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>")
    body.append("</table>")

    css = ("* {font-family: sans-serif; font-size: 10pt;} h1 {font-size: 16pt;} h2 {font-size: 12pt; margin-top: 12pt;} "
           "h3 {font-size: 10.5pt; margin-top: 6pt;} p {margin-top: 3pt; text-align: justify;} "
           "td, th {border: 0.5pt solid gray; padding: 3pt;}")
    story = pymupdf.Story(html="\n".join(body), user_css=css)
    page_rect = pymupdf.paper_rect("a4")
    content = page_rect + (60, 60, -60, -60)
    tmp = path.with_suffix(".tmp.pdf")
    writer = pymupdf.DocumentWriter(str(tmp))
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content)
        story.draw(device)
        writer.end_page()
    writer.close()

    # Running header and footer, as most real policy PDFs have; the parser should drop them.
    with pymupdf.open(tmp) as doc:
        for number, page in enumerate(doc, 1):
            page.insert_text((60, 40), f"{ISSUER} - {TITLE}", fontsize=8, color=(0.4, 0.4, 0.4))
            page.insert_text((60, page_rect.height - 30), f"Page {number} of {len(doc)}", fontsize=8,
                             color=(0.4, 0.4, 0.4))
        doc.save(path)
    tmp.unlink()


def write_docx(path: Path) -> None:
    import docx

    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_heading(TITLE, level=0)
    document.add_paragraph(ISSUER)
    document.add_paragraph(VERSION)
    for level, heading, paragraphs in SECTIONS:
        document.add_heading(heading, level=level)
        for p in paragraphs:
            document.add_paragraph(p)
    title, header, rows = ANNEX_A
    document.add_heading(title, level=1)
    table = document.add_table(rows=1, cols=len(header))
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    document.save(str(path))
